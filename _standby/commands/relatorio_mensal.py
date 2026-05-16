"""
relatorio_mensal.py — Relatório completo de performance mensal.
Uso: python sbwaa.py /relatorio-mensal
"""

import re
import sys
import json
from datetime import datetime, timedelta, date
from pathlib import Path

import yfinance as yf
import pandas as pd
from docx import Document

PROJECT_ROOT = Path(__file__).parent.parent.parent
VAULT_ROOT = PROJECT_ROOT / "vault"
SCRIPTS_DATA = PROJECT_ROOT / "scripts" / "data"
CACHE_DIR = SCRIPTS_DATA / "cache"
CARTEIRA_PATH = VAULT_ROOT / "00-portfolio" / "carteira.md"
MENSAIS_DIR = VAULT_ROOT / "02-relatorios" / "mensais"
IPS_PATH = VAULT_ROOT / "00-portfolio" / "ips.md"

SELIC_PADRAO = 0.1275  # CDI aproximado


def eh_ticker_br(ticker: str) -> bool:
    return bool(re.match(r"^[A-Z]{4}[0-9]{1,2}[FBP]?$", ticker))


def ticker_yahoo(ticker: str) -> str:
    return f"{ticker}.SA" if eh_ticker_br(ticker) else ticker


def parse_carteira() -> list[dict]:
    if not CARTEIRA_PATH.exists():
        return []
    conteudo = CARTEIRA_PATH.read_text(encoding="utf-8")
    posicoes = []
    dentro = False
    for linha in conteudo.splitlines():
        stripped = linha.strip()
        if stripped.startswith("| Ticker"):
            dentro = True
            continue
        if dentro and stripped.startswith("|---"):
            continue
        if dentro and stripped.startswith("|"):
            cols = [c.strip() for c in stripped.split("|")[1:-1]]
            if len(cols) >= 5 and cols[0] and cols[0] not in ("", "Ticker"):
                try:
                    qtd_s = cols[3].replace(".", "").replace(",", ".")
                    pm_s = re.sub(r"[^\d,.]", "", cols[4]).replace(",", ".")
                    qtd = float(qtd_s) if qtd_s else 0
                    pm = float(pm_s) if pm_s else 0
                    if cols[0] and qtd > 0:
                        posicoes.append({"ticker": cols[0], "tipo": cols[1], "qtd": qtd, "pm": pm})
                except (ValueError, IndexError):
                    pass
        elif dentro and not stripped.startswith("|"):
            break
    return posicoes


def calcular_performance_mes(posicoes: list[dict], ini: date, fim: date) -> list[dict]:
    resultados = []
    for pos in posicoes:
        try:
            ty = ticker_yahoo(pos["ticker"])
            hist = yf.download(ty, start=str(ini - timedelta(days=5)), end=str(fim + timedelta(days=1)),
                               auto_adjust=True, progress=False, threads=False)
            if hist.empty:
                continue
            close = hist["Close"].squeeze()
            pa_ini = float(close.iloc[0])
            pa_fim = float(close.iloc[-1])
            var_pct = (pa_fim / pa_ini - 1) * 100
            pl_mes = (pa_fim - pa_ini) * pos["qtd"]
            valor_atual = pa_fim * pos["qtd"]
            valor_ini = pa_ini * pos["qtd"]
            resultados.append({
                "ticker": pos["ticker"],
                "tipo": pos["tipo"],
                "preco_ini": pa_ini,
                "preco_fim": pa_fim,
                "var_pct": var_pct,
                "pl_mes": pl_mes,
                "valor_atual": valor_atual,
                "valor_ini": valor_ini,
            })
        except Exception:
            pass
    return resultados


def benchmark_mes(symbol: str, ini: date, fim: date) -> float | None:
    try:
        hist = yf.download(symbol, start=str(ini - timedelta(days=5)), end=str(fim + timedelta(days=1)),
                           auto_adjust=True, progress=False, threads=False)
        if hist.empty:
            return None
        close = hist["Close"].squeeze()
        return float((close.iloc[-1] / close.iloc[0] - 1) * 100)
    except Exception:
        return None


def gerar_docx(conteudo: str, out_path: Path):
    doc = Document()
    doc.add_heading("SBWAA — Relatório Mensal", 0)
    for linha in conteudo.splitlines():
        if linha.startswith("## "):
            doc.add_heading(linha[3:], 2)
        elif linha.startswith("# "):
            doc.add_heading(linha[2:], 1)
        elif linha.strip():
            doc.add_paragraph(linha)
    doc.save(str(out_path))


def main():
    hoje = date.today()
    ini_mes = date(hoje.year, hoje.month, 1)
    nome_mes = hoje.strftime("%Y-%m")
    nome_mes_ext = hoje.strftime("%B/%Y")

    print(f"\n{'═'*55}")
    print(f"  SBWAA — Relatório Mensal | {nome_mes_ext}")
    print(f"{'═'*55}\n")

    posicoes = parse_carteira()
    if not posicoes:
        print("  Carteira vazia.\n")
        return

    print("Calculando performance do mês (pode levar ~1 min)...")
    resultados = calcular_performance_mes(posicoes, ini_mes, hoje)

    ibov = benchmark_mes("^BVSP", ini_mes, hoje)
    sp500 = benchmark_mes("^GSPC", ini_mes, hoje)

    # CDI do mês (aproximado)
    dias_uteis = (hoje - ini_mes).days * 5 / 7
    cdi_mes = ((1 + SELIC_PADRAO) ** (dias_uteis / 252) - 1) * 100

    total_pl = sum(r["pl_mes"] for r in resultados)
    total_ini = sum(r["valor_ini"] for r in resultados)
    total_fim = sum(r["valor_atual"] for r in resultados)
    carteira_pct = (total_fim / total_ini - 1) * 100 if total_ini > 0 else 0

    tabela = "\n".join(
        f"| {r['ticker']:<8} | {r['tipo']:<12} | {r['var_pct']:>+7.2f}% | R$ {r['pl_mes']:>+10,.2f} | "
        f"{r['valor_atual'] / total_fim * 100 if total_fim > 0 else 0:.1f}% |"
        for r in sorted(resultados, key=lambda x: x["var_pct"], reverse=True)
    )

    # Dividendos do mês
    divs_mes = {}
    for pos in posicoes:
        try:
            t = yf.Ticker(ticker_yahoo(pos["ticker"]))
            divs = t.dividends
            if divs is not None and not divs.empty:
                if hasattr(divs.index, "tz"):
                    divs.index = divs.index.tz_localize(None) if divs.index.tzinfo else divs.index
                d_mes = divs[(divs.index >= pd.Timestamp(ini_mes)) & (divs.index <= pd.Timestamp(hoje))]
                if not d_mes.empty:
                    divs_mes[pos["ticker"]] = float(d_mes.sum()) * pos["qtd"]
        except Exception:
            pass
    total_divs = sum(divs_mes.values())

    # Decisões do PM no mês
    decisoes_md = ""
    dec_path = VAULT_ROOT / "00-portfolio" / "decisoes.md"
    if dec_path.exists():
        for linha in dec_path.read_text(encoding="utf-8").splitlines():
            if nome_mes in linha and "|" in linha:
                decisoes_md += f"  {linha}\n"

    import anthropic
    client = anthropic.Anthropic()

    prompt = f"""DATA: {hoje.strftime('%Y-%m-%d')} | Mês: {nome_mes_ext}

PERFORMANCE DO MÊS:
| Ticker | Tipo | Var% | P&L Mês | Aloc% |
|--------|------|------|---------|-------|
{tabela}

BENCHMARKS:
- Carteira: {carteira_pct:+.2f}%
- IBOV: {f'{ibov:+.2f}%' if ibov else 'N/D'}
- S&P 500: {f'{sp500:+.2f}%' if sp500 else 'N/D'}
- CDI (aprox): {cdi_mes:+.2f}%

PATRIMÔNIO:
- Início do mês: R$ {total_ini:,.2f}
- Fim (atual): R$ {total_fim:,.2f}
- P&L Mês: R$ {total_pl:+,.2f}

PROVENTOS DO MÊS: R$ {total_divs:,.2f}
{json.dumps(divs_mes, ensure_ascii=False) if divs_mes else 'Sem dividendos registrados.'}

DECISÕES DO PM NO MÊS:
{decisoes_md if decisoes_md else 'Sem decisões registradas.'}

Gere o Relatório Mensal completo:

---
tags: [relatorio, mensal]
cssclasses: [node-relatorio]
mes: {nome_mes}
---

# Relatório Mensal — {nome_mes_ext}

## 📊 Performance
{{tabela completa com atribuição de retorno}}

## 🏆 Destaques
{{melhor e pior ativo, comparação vs benchmarks}}

## 💰 Resumo Financeiro
{{patrimônio início vs fim, P&L, dividendos}}

## 📈 Atribuição de Retorno
{{quanto cada ativo contribuiu para o resultado total}}

## 📉 Evolução de Risco
{{Sharpe, VaR e drawdown ao longo do mês — se dados disponíveis}}

## 🔄 Decisões do PM
{{revisão das decisões tomadas no mês}}

## 📅 Próximos Eventos
{{top 3-5 eventos relevantes para o próximo mês: earnings, vencimentos, COPOM}}

## Links
[[carteira]] | [[ips]] | [[decisoes]]
"""

    print("Gerando relatório mensal (claude-sonnet-4-6)...")
    response = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=2500,
        messages=[{"role": "user", "content": prompt}],
    )
    output_md = response.content[0].text

    MENSAIS_DIR.mkdir(parents=True, exist_ok=True)
    md_path = MENSAIS_DIR / f"relatorio-{nome_mes}.md"
    docx_path = MENSAIS_DIR / f"relatorio-{nome_mes}.docx"
    md_path.write_text(output_md, encoding="utf-8")
    gerar_docx(output_md, docx_path)

    print(output_md[:400] + "...\n")
    print(f"{'═'*55}")
    print(f"  ✓ {md_path.relative_to(PROJECT_ROOT)}")
    print(f"  ✓ {docx_path.relative_to(PROJECT_ROOT)}")
    print(f"{'═'*55}\n")


if __name__ == "__main__":
    main()
