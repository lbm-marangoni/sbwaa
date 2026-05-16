"""
relatorio_semanal.py — Relatório de performance da semana.
Uso: python sbwaa.py /relatorio-semanal
"""

import re
import sys
import math
import json
from datetime import datetime, timedelta, date
from pathlib import Path

import yfinance as yf
import pandas as pd
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_ROOT = Path(__file__).parent.parent.parent
VAULT_ROOT = PROJECT_ROOT / "vault"
SCRIPTS_DATA = PROJECT_ROOT / "scripts" / "data"
CACHE_DIR = SCRIPTS_DATA / "cache"
CARTEIRA_PATH = VAULT_ROOT / "00-portfolio" / "carteira.md"
SEMANAIS_DIR = VAULT_ROOT / "02-relatorios" / "semanais"


def eh_ticker_br(ticker: str) -> bool:
    return bool(re.match(r"^[A-Z]{4}[0-9]{1,2}[FBP]?$", ticker))


def ticker_yahoo(ticker: str) -> str:
    return f"{ticker}.SA" if eh_ticker_br(ticker) else ticker


def parse_carteira() -> list[dict]:
    if not CARTEIRA_PATH.exists():
        return []
    conteudo = CARTEIRA_PATH.read_text(encoding="utf-8")
    posicoes = []
    dentro = False
    for linha in conteudo.splitlines():
        stripped = linha.strip()
        if stripped.startswith("| Ticker"):
            dentro = True
            continue
        if dentro and stripped.startswith("|---"):
            continue
        if dentro and stripped.startswith("|"):
            cols = [c.strip() for c in stripped.split("|")[1:-1]]
            if len(cols) >= 5 and cols[0] and cols[0] not in ("", "Ticker"):
                try:
                    qtd_s = cols[3].replace(".", "").replace(",", ".")
                    pm_s = re.sub(r"[^\d,.]", "", cols[4]).replace(",", ".")
                    qtd = float(qtd_s) if qtd_s else 0
                    pm = float(pm_s) if pm_s else 0
                    if cols[0] and qtd > 0:
                        posicoes.append({"ticker": cols[0], "tipo": cols[1], "qtd": qtd, "pm": pm})
                except (ValueError, IndexError):
                    pass
        elif dentro and not stripped.startswith("|"):
            break
    return posicoes


def calcular_performance_semana(posicoes: list[dict]) -> list[dict]:
    hoje = date.today()
    seg = hoje - timedelta(days=hoje.weekday())  # Segunda
    sex = seg + timedelta(days=4)  # Sexta

    resultados = []
    for pos in posicoes:
        try:
            ty = ticker_yahoo(pos["ticker"])
            hist = yf.download(ty, start=str(seg - timedelta(days=3)), end=str(hoje + timedelta(days=1)),
                               auto_adjust=True, progress=False, threads=False)
            if hist.empty:
                continue
            close = hist["Close"].squeeze()
            pa_ini = float(close.iloc[0])
            pa_fim = float(close.iloc[-1])
            var_pct = (pa_fim / pa_ini - 1) * 100
            pl_semana = (pa_fim - pa_ini) * pos["qtd"]
            valor_atual = pa_fim * pos["qtd"]
            resultados.append({
                "ticker": pos["ticker"],
                "tipo": pos["tipo"],
                "preco_ini": pa_ini,
                "preco_fim": pa_fim,
                "var_pct": var_pct,
                "pl_semana": pl_semana,
                "valor_atual": valor_atual,
            })
        except Exception:
            pass
    return resultados


def carregar_ibov_semana() -> float | None:
    try:
        hoje = date.today()
        seg = hoje - timedelta(days=hoje.weekday())
        hist = yf.download("^BVSP", start=str(seg - timedelta(days=3)),
                           end=str(hoje + timedelta(days=1)),
                           auto_adjust=True, progress=False, threads=False)
        if hist.empty:
            return None
        close = hist["Close"].squeeze()
        return float((close.iloc[-1] / close.iloc[0] - 1) * 100)
    except Exception:
        return None


def gerar_docx(md_conteudo: str, out_path: Path):
    doc = Document()
    doc.add_heading("SBWAA — Relatório Semanal", 0)
    for linha in md_conteudo.splitlines():
        if linha.startswith("## "):
            doc.add_heading(linha[3:], 2)
        elif linha.startswith("# "):
            doc.add_heading(linha[2:], 1)
        elif linha.startswith("| "):
            # Tabela simples
            cols = [c.strip() for c in linha.split("|")[1:-1]]
            if cols:
                table = doc.add_table(rows=1, cols=len(cols))
                table.style = "Table Grid"
                for i, c in enumerate(cols):
                    table.rows[0].cells[i].text = c
        elif linha.strip():
            doc.add_paragraph(linha)
    doc.save(str(out_path))


def main():
    hoje = date.today()
    num_semana = hoje.isocalendar()[1]
    ano = hoje.year
    nome = f"semana-{ano}-W{num_semana:02d}"

    print(f"\n{'═'*55}")
    print(f"  SBWAA — Relatório Semanal | Semana {num_semana}/{ano}")
    print(f"{'═'*55}\n")

    posicoes = parse_carteira()
    if not posicoes:
        print("  Carteira vazia.\n")
        return

    print("Calculando performance da semana...")
    resultados = calcular_performance_semana(posicoes)
    ibov_semana = carregar_ibov_semana()

    if not resultados:
        print("⚠️  Sem dados de preços disponíveis.\n")
        return

    total_pl = sum(r["pl_semana"] for r in resultados)
    total_valor = sum(r["valor_atual"] for r in resultados)
    melhor = max(resultados, key=lambda x: x["var_pct"])
    pior = min(resultados, key=lambda x: x["var_pct"])

    # Tabela de performance
    tabela_linhas = "\n".join(
        f"| {r['ticker']:<8} | {r['tipo']:<12} | R$ {r['preco_ini']:>8,.2f} | R$ {r['preco_fim']:>8,.2f} | {r['var_pct']:>+7.2f}% | R$ {r['pl_semana']:>+10,.2f} |"
        for r in sorted(resultados, key=lambda x: x["var_pct"], reverse=True)
    )

    ibov_txt = f"{ibov_semana:+.2f}%" if ibov_semana is not None else "N/D"

    import anthropic
    client = anthropic.Anthropic()

    prompt = f"""DATA: {hoje.strftime('%Y-%m-%d')} | Semana {num_semana}/{ano}

PERFORMANCE DA SEMANA:
| Ticker | Tipo | P.Ini | P.Fim | Var% | P&L Semana |
|--------|------|-------|-------|------|------------|
{tabela_linhas}

RESUMO:
- P&L Total da Semana: R$ {total_pl:+,.2f}
- Patrimônio Atual Estimado: R$ {total_valor:,.2f}
- Melhor da semana: {melhor['ticker']} ({melhor['var_pct']:+.2f}%)
- Pior da semana: {pior['ticker']} ({pior['var_pct']:+.2f}%)
- IBOV na semana: {ibov_txt}

Gere o Relatório Semanal completo com este formato:

---
tags: [relatorio, semanal]
cssclasses: [node-relatorio]
semana: {num_semana}
ano: {ano}
---

# Relatório Semanal — Semana {num_semana}/{ano}

## 📊 Performance da Semana

{{tabela com todos os ativos}}

## 🏆 Destaques
- Melhor: {melhor['ticker']} ({melhor['var_pct']:+.2f}%)
- Pior: {pior['ticker']} ({pior['var_pct']:+.2f}%)
- vs IBOV: {{carteira vs {ibov_txt}}}

## 💰 Resumo Financeiro
{{patrimônio início vs fim, P&L total, variação %}}

## 📈 Análise
{{2-3 parágrafos: o que moveu a carteira, principais fatores, contexto macro}}

## 🔭 Próxima Semana
{{top 3 eventos a monitorar}}

## Links
{{links dos morning-calls da semana}}
"""

    print("Gerando relatório (claude-sonnet-4-6)...")
    response = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=2000,
        messages=[{"role": "user", "content": prompt}],
    )
    output_md = response.content[0].text

    # Salvar MD e DOCX
    SEMANAIS_DIR.mkdir(parents=True, exist_ok=True)
    md_path = SEMANAIS_DIR / f"{nome}.md"
    docx_path = SEMANAIS_DIR / f"{nome}.docx"

    md_path.write_text(output_md, encoding="utf-8")
    gerar_docx(output_md, docx_path)

    print(output_md[:500] + "...\n")
    print(f"{'═'*55}")
    print(f"  ✓ {md_path.relative_to(PROJECT_ROOT)}")
    print(f"  ✓ {docx_path.relative_to(PROJECT_ROOT)}")
    print(f"{'═'*55}\n")


if __name__ == "__main__":
    main()
