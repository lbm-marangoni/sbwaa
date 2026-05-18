"""
run_valuation_reviewer.py — Executa o agente Valuation Reviewer do SBWAA.
Uso: python run_valuation_reviewer.py PETR4 [--versao curta|longa]
"""

import sys
import json
import argparse
import subprocess
from datetime import datetime, date
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_ROOT = Path(__file__).parent.parent.parent.parent
VAULT_ROOT = PROJECT_ROOT / "vault"
SCRIPTS_DATA = PROJECT_ROOT / "scripts" / "data"
AGENTS_DIR = PROJECT_ROOT / ".claude" / "agents"
SKILL_PATH = Path(__file__).parent / "SKILL.md"
CACHE_DIR = SCRIPTS_DATA / "cache"

sys.path.insert(0, str(SCRIPTS_DATA))


# ─── Pixel Art ────────────────────────────────────────────────────────────────

def _inserir_pixel_art_docx(doc: Document, agente_id: str):
    assets_dir = PROJECT_ROOT / "vault" / "assets" / "agents-pixel"
    img_path = assets_dir / f"{agente_id}.png"
    if not img_path.exists():
        return
    section = doc.sections[0]
    header = section.header
    para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    run = para.add_run()
    run.add_picture(str(img_path), width=Inches(0.5))


# ─── Loaders ──────────────────────────────────────────────────────────────────

def carregar_dcf_cache(ticker: str, hoje: str) -> dict | None:
    path = CACHE_DIR / f"dcf_{ticker}_{hoje}.json"
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8"))
    # Tentar data anterior (até 7 dias)
    for d in range(1, 8):
        from datetime import timedelta
        dt = (datetime.strptime(hoje, "%Y-%m-%d") - timedelta(days=d)).strftime("%Y-%m-%d")
        path = CACHE_DIR / f"dcf_{ticker}_{dt}.json"
        if path.exists():
            print(f"DCF encontrado do dia {dt}.")
            return json.loads(path.read_text(encoding="utf-8"))
    return None


def carregar_earnings_mais_recente(ticker: str) -> str:
    ativo_dir = VAULT_ROOT / "01-ativos" / ticker
    if not ativo_dir.exists():
        return ""
    arquivos = sorted(ativo_dir.glob("earnings-*.md"), reverse=True)
    return arquivos[0].read_text(encoding="utf-8")[:2000] if arquivos else ""


def carregar_market_researcher_hoje(hoje: str) -> str:
    path = VAULT_ROOT / "03-macro" / f"market-researcher-{hoje}.md"
    return path.read_text(encoding="utf-8")[:1500] if path.exists() else ""


# ─── DOCX Builder ─────────────────────────────────────────────────────────────

def _add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)


def _add_hr(doc: Document):
    doc.add_paragraph("─" * 60)


def gerar_docx(ticker: str, hoje: str, conteudo_md: str, versao: str, saida: Path) -> None:
    doc = Document()

    # Margens
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Cabeçalho
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.text = f"SBWAA — Equity Research | {ticker} | {hoje}"
    hp.runs[0].font.size = Pt(9)
    hp.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Rodapé
    footer = doc.sections[0].footer
    fp = footer.paragraphs[0]
    fp.text = f"SBWAA — Confidencial | Gerado em {hoje}"
    fp.runs[0].font.size = Pt(8)
    fp.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Pixel art do agente no cabeçalho
    _inserir_pixel_art_docx(doc, "valuation-reviewer")

    # Título principal
    titulo = doc.add_heading(f"Equity Research — {ticker}", 0)
    titulo.runs[0].font.size = Pt(16)
    titulo.runs[0].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    sub = doc.add_paragraph(f"Valuation Reviewer SBWAA | {versao.upper()} | {hoje}")
    sub.runs[0].font.size = Pt(9)
    sub.runs[0].font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    _add_hr(doc)

    # Corpo: parsear markdown básico → DOCX
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            return
        # Filtrar linhas separadoras (|---|)
        data_rows = [r for r in table_rows if not all(c.strip().startswith("-") for c in r)]
        if not data_rows:
            table_rows = []
            in_table = False
            return
        cols = len(data_rows[0])
        t = doc.add_table(rows=len(data_rows), cols=cols)
        t.style = "Table Grid"
        for ri, row in enumerate(data_rows):
            for ci, cell_text in enumerate(row):
                cell = t.cell(ri, ci)
                cell.text = cell_text.strip()
                if ri == 0:
                    cell.paragraphs[0].runs[0].bold = True if cell.paragraphs[0].runs else True
        doc.add_paragraph("")
        table_rows = []
        in_table = False

    for linha in conteudo_md.splitlines():
        stripped = linha.strip()

        # Frontmatter — pular
        if stripped.startswith("---"):
            continue
        if stripped.startswith("tags:") or stripped.startswith("cssclasses:") \
                or stripped.startswith("data:") or stripped.startswith("versao:") \
                or stripped.startswith("agente:") or stripped.startswith("ticker:"):
            continue

        # Tabela
        if stripped.startswith("|"):
            cells = [c for c in stripped.split("|")[1:-1]]
            if in_table:
                table_rows.append(cells)
            else:
                in_table = True
                table_rows.append(cells)
            continue
        else:
            if in_table:
                flush_table()

        if not stripped:
            doc.add_paragraph("")
            continue

        if stripped.startswith("# "):
            _add_heading(doc, stripped[2:], 1)
        elif stripped.startswith("## "):
            _add_heading(doc, stripped[3:], 2)
        elif stripped.startswith("### "):
            _add_heading(doc, stripped[4:], 3)
        elif stripped.startswith("**") and stripped.endswith("**") and len(stripped) > 4:
            p = doc.add_paragraph()
            run = p.add_run(stripped[2:-2])
            run.bold = True
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(stripped[2:], style="List Bullet")
        else:
            # Texto com negrito inline
            p = doc.add_paragraph()
            partes = stripped.split("**")
            for i, parte in enumerate(partes):
                run = p.add_run(parte)
                if i % 2 == 1:
                    run.bold = True

    if in_table:
        flush_table()

    doc.save(saida)


# ─── Atualizar tese.md ────────────────────────────────────────────────────────

def atualizar_tese_com_link(ticker: str, nome_arquivo: str) -> None:
    tese_path = VAULT_ROOT / "01-ativos" / ticker / "tese.md"
    if not tese_path.exists():
        return
    conteudo = tese_path.read_text(encoding="utf-8")
    link = f"- [[{ticker}/{nome_arquivo}]]"
    if link not in conteudo:
        if "## Valuation" in conteudo:
            conteudo = conteudo.replace("## Valuation", f"## Valuation\n{link}", 1)
        else:
            conteudo += f"\n\n## Valuation\n{link}\n"
        tese_path.write_text(conteudo, encoding="utf-8")
        print(f"Wikilink adicionado à tese: {tese_path}")


# ─── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Valuation Reviewer — SBWAA")
    parser.add_argument("ticker", help="Ticker do ativo (ex: PETR4)")
    parser.add_argument("--versao", choices=["curta", "longa"], default="curta",
                        help="Versão do relatório (default: curta)")
    args = parser.parse_args()

    ticker = args.ticker.upper().strip()
    versao = args.versao
    hoje = datetime.now().strftime("%Y-%m-%d")

    print(f"Valuation Reviewer — {ticker} | versão {versao} | {hoje}")
    print("=" * 55)

    # Carregar DCF ou rodar Model Builder
    dcf = carregar_dcf_cache(ticker, hoje)
    if dcf is None:
        print("DCF não encontrado no cache. Rodando Model Builder...")
        resultado = subprocess.run(
            [sys.executable,
             str(AGENTS_DIR / "model-builder" / "run_model_builder.py"),
             ticker],
            text=True,
        )
        if resultado.returncode != 0:
            print("ERRO ao rodar Model Builder. Abortando.")
            sys.exit(1)
        dcf = carregar_dcf_cache(ticker, hoje)
        if dcf is None:
            print("ERRO: DCF ainda não disponível após Model Builder.")
            sys.exit(1)

    from fetch_brapi import buscar_ticker as brapi_buscar
    try:
        dados_brapi = brapi_buscar(ticker)
    except Exception as e:
        print(f"AVISO: Erro ao buscar Brapi para {ticker}: {e}")
        dados_brapi = {}

    earnings_ctx = carregar_earnings_mais_recente(ticker)
    macro_ctx = carregar_market_researcher_hoje(hoje)
    skill_content = SKILL_PATH.read_text(encoding="utf-8")

    res = dcf.get("resultado", {})
    prem = dcf.get("premissas", {})

    contexto = f"""TICKER: {ticker}
DATA: {hoje}
VERSÃO SOLICITADA: {versao.upper()}

DCF (Model Builder):
- WACC: {prem.get('wacc', 0)*100:.2f}%
- g perpetuidade: {prem.get('g_perpetuidade', 0)*100:.2f}%
- Valor Justo: R$ {res.get('valor_justo', 0):.2f}
- Cotação Atual: R$ {res.get('cotacao_atual', 0):.2f}
- Upside: {res.get('upside_pct', 0):+.1f}%
- Enterprise Value: {res.get('enterprise_value', 'N/D')}
- Equity Value: {res.get('equity_value', 'N/D')}
- Notas do modelo: {dcf.get('notas', [])}

SENSIBILIDADE DCF:
{json.dumps(dcf.get('sensibilidade', {}), indent=2)}

MÚLTIPLOS BRAPI (atuais):
- P/L: {dados_brapi.get('pl', 'N/D')}
- EV/EBITDA: {dados_brapi.get('ev_ebitda', 'N/D')}
- P/VP: {dados_brapi.get('pvp', 'N/D')}
- DY: {dados_brapi.get('dy', 'N/D')}%
- ROE: {dados_brapi.get('roe', 'N/D')}%
- Setor: {dados_brapi.get('setor', 'N/D')}
- Nome: {dados_brapi.get('nome', 'N/D')}

{f"EARNINGS REVIEWER:{chr(10)}{earnings_ctx}" if earnings_ctx else "EARNINGS REVIEWER: Não disponível."}

{f"MARKET RESEARCHER:{chr(10)}{macro_ctx}" if macro_ctx else "MARKET RESEARCHER: Não disponível para hoje."}

Gere o relatório na versão {versao.upper()}, com o formato exato definido no SKILL.
"""

    # RAG: buscar benchmarks de valuation e metodologia do setor
    sys.path.insert(0, str(PROJECT_ROOT))
    contexto_rag = ""
    try:
        from knowledge.retriever import buscar, formatar_contexto_para_agente, base_disponivel
        if base_disponivel():
            setor_val = dados_brapi.get("setor", "")
            query_val = f"valuation múltiplos {ticker} {setor_val} DCF equity research"
            chunks = buscar(query_val, n_resultados=4, filtro_tipo=None)
            contexto_rag = formatar_contexto_para_agente(chunks, max_tokens=1500)
    except Exception:
        pass

    if contexto_rag:
        contexto += f"\n\n{contexto_rag}"

    import anthropic
    print(f"Enviando ao Valuation Reviewer (claude-sonnet-4-6) — versão {versao}...")
    client = anthropic.Anthropic()
    response = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=2000,
        system=skill_content,
        messages=[{"role": "user", "content": contexto}],
    )

    output_md = response.content[0].text

    ativo_dir = VAULT_ROOT / "01-ativos" / ticker
    ativo_dir.mkdir(parents=True, exist_ok=True)

    nome_base = f"equity-research-{ticker}-{hoje}-{versao}"
    md_path = ativo_dir / f"{nome_base}.md"
    md_path.write_text(output_md, encoding="utf-8")
    print(f"Markdown salvo: {md_path}")

    docx_path = ativo_dir / f"{nome_base}.docx"
    gerar_docx(ticker, hoje, output_md, versao, docx_path)
    print(f"DOCX gerado: {docx_path}")

    atualizar_tese_com_link(ticker, nome_base)

    # Extrair e exibir veredicto
    veredicto = "N/D"
    upside_str = f"{res.get('upside_pct', 0):+.1f}%"
    for linha in output_md.splitlines():
        if "VEREDICTO" in linha.upper() and any(v in linha.upper() for v in ["BARATO", "JUSTO", "CARO"]):
            veredicto = linha.strip()
            break

    print(f"\n{'═'*50}")
    print(f"  {ticker} — Upside DCF: {upside_str}")
    print(f"  {veredicto}")
    print(f"{'═'*50}\n")


if __name__ == "__main__":
    main()
