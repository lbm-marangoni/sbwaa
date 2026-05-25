"""
run_model_builder.py — Executa o agente Model Builder (DCF) do SBWAA.
Uso: python run_model_builder.py PETR4 [--tipo fii|etf]
"""

import sys
import json
import re
import argparse
import subprocess
from datetime import datetime, date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).parent.parent.parent.parent
VAULT_ROOT = PROJECT_ROOT / "vault"
SCRIPTS_DATA = PROJECT_ROOT / "scripts" / "data"
SKILL_PATH = Path(__file__).parent / "SKILL.md"
CACHE_DIR = SCRIPTS_DATA / "cache"

sys.path.insert(0, str(SCRIPTS_DATA))


# ─── Helpers ──────────────────────────────────────────────────────────────────

def fmt_pct(v): return f"{v*100:.2f}%" if v is not None else "N/D"
def fmt_brl(v): return f"R$ {v:,.0f}" if v is not None else "N/D"
def fmt_2(v): return f"{v:.2f}" if v is not None else "N/D"


def extrair_json(texto: str) -> dict:
    texto = texto.strip()
    # Remover code fences se presentes
    texto = re.sub(r"^```(?:json)?\s*", "", texto, flags=re.MULTILINE)
    texto = re.sub(r"```\s*$", "", texto, flags=re.MULTILINE)
    # Encontrar o bloco JSON
    match = re.search(r"\{.*\}", texto, re.DOTALL)
    if not match:
        raise ValueError("Nenhum JSON encontrado na resposta do modelo.")
    return json.loads(match.group())


def carregar_earnings_mais_recente(ticker: str) -> str:
    ativo_dir = VAULT_ROOT / "01-ativos" / ticker
    if not ativo_dir.exists():
        return ""
    arquivos = sorted(ativo_dir.glob(f"earnings-*.md"), reverse=True)
    if arquivos:
        return arquivos[0].read_text(encoding="utf-8")[:2000]
    return ""


def carregar_market_researcher_hoje(hoje: str) -> str:
    path = VAULT_ROOT / "03-macro" / f"market-researcher-{hoje}.md"
    if path.exists():
        return path.read_text(encoding="utf-8")[:1500]
    return ""


# ─── XLSX Builder ─────────────────────────────────────────────────────────────

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ACCENT_FILL = PatternFill("solid", fgColor="D6E4F0")
BORDER_THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


def _hdr(ws, row, col, value):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = HEADER_FILL
    c.font = HEADER_FONT
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER_THIN


def _cel(ws, row, col, value, bold=False, accent=False):
    c = ws.cell(row=row, column=col, value=value)
    c.border = BORDER_THIN
    if bold:
        c.font = Font(bold=True)
    if accent:
        c.fill = ACCENT_FILL


def criar_xlsx(dcf: dict, ticker: str, saida: Path) -> None:
    wb = openpyxl.Workbook()
    p = dcf.get("premissas", {})
    proj = dcf.get("projecoes", {})
    res = dcf.get("resultado", {})
    vt = dcf.get("valor_terminal", {})
    sens = dcf.get("sensibilidade", {})
    notas = dcf.get("notas", [])

    # ── Aba 1: Premissas ──────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Premissas"
    ws1.column_dimensions["A"].width = 32
    ws1.column_dimensions["B"].width = 18
    ws1.column_dimensions["C"].width = 30

    _hdr(ws1, 1, 1, "Premissa")
    _hdr(ws1, 1, 2, "Valor")
    _hdr(ws1, 1, 3, "Fonte / Observação")

    premissas_rows = [
        ("Taxa livre de risco (Rf)", fmt_pct(p.get("risk_free")), "Selic / NTN-B 10a"),
        ("Prêmio de risco mercado", fmt_pct(p.get("premio_risco")), "Padrão SBWAA"),
        ("Beta", fmt_2(p.get("beta")), "Brapi / estimado por setor"),
        ("Custo equity (Ke)", fmt_pct(p.get("ke")), "CAPM"),
        ("Custo dívida (Kd)", fmt_pct(p.get("kd")), "Estimado"),
        ("Peso equity", fmt_pct(p.get("peso_equity")), "Balanço"),
        ("Peso dívida", fmt_pct(p.get("peso_divida")), "Balanço"),
        ("IR efetivo", fmt_pct(p.get("ir_efetivo")), "Padrão 34%"),
        ("WACC", fmt_pct(p.get("wacc")), "CAPM ponderado"),
        ("g perpetuidade", fmt_pct(p.get("g_perpetuidade")), "PIB BR nominal LP"),
        ("Margem EBITDA base", fmt_pct(p.get("margem_ebitda_base")), "Média histórica"),
        ("Capex / Receita", fmt_pct(p.get("capex_pct_receita")), "Histórico"),
        ("ΔGiro / ΔReceita", fmt_pct(p.get("giro_capital_pct")), "Estimado"),
        ("Crescimento receita Ano 1", fmt_pct(p.get("crescimento_receita_ano1")), ""),
        ("Crescimento receita Ano 2", fmt_pct(p.get("crescimento_receita_ano2")), ""),
        ("Crescimento receita Ano 3", fmt_pct(p.get("crescimento_receita_ano3")), ""),
        ("Crescimento receita Ano 4", fmt_pct(p.get("crescimento_receita_ano4")), ""),
        ("Crescimento receita Ano 5", fmt_pct(p.get("crescimento_receita_ano5")), ""),
    ]
    for i, (nome, val, fonte) in enumerate(premissas_rows, start=2):
        _cel(ws1, i, 1, nome)
        _cel(ws1, i, 2, val, accent=(i % 2 == 0))
        _cel(ws1, i, 3, fonte)

    if notas:
        r = len(premissas_rows) + 3
        ws1.cell(row=r, column=1, value="NOTAS DO MODELO:").font = Font(bold=True)
        for n in notas:
            r += 1
            ws1.cell(row=r, column=1, value=f"• {n}")

    # ── Aba 2: Projeções ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Projeções")
    ws2.column_dimensions["A"].width = 26
    for col in range(2, 8):
        ws2.column_dimensions[get_column_letter(col)].width = 16

    headers = ["Métrica", "Ano 1", "Ano 2", "Ano 3", "Ano 4", "Ano 5"]
    for col, h in enumerate(headers, start=1):
        _hdr(ws2, 1, col, h)

    anos = [proj.get(f"ano{i}", {}) for i in range(1, 6)]
    metricas = [
        ("Receita Líquida (R$ MM)", "receita", True),
        ("EBITDA (R$ MM)", "ebitda", False),
        ("Margem EBITDA", "ebitda_margem", False),
        ("Capex (R$ MM)", "capex", False),
        ("ΔCapital de Giro (R$ MM)", "delta_giro", False),
        ("FCFF (R$ MM)", "fcff", True),
        ("FCFF Descontado (R$ MM)", "fcff_descontado", True),
    ]
    for row_i, (label, key, bold) in enumerate(metricas, start=2):
        _cel(ws2, row_i, 1, label, bold=bold)
        for col_i, ano in enumerate(anos, start=2):
            val = ano.get(key, 0)
            if key == "ebitda_margem":
                val = fmt_pct(val) if isinstance(val, float) else val
            _cel(ws2, row_i, col_i, val, bold=bold, accent=(row_i % 2 == 0))

    # Valor terminal
    r_vt = len(metricas) + 3
    ws2.cell(row=r_vt, column=1, value="VALOR TERMINAL").font = Font(bold=True, color="1F3864")
    vt_rows = [
        ("FCFF Normalizado (R$ MM)", vt.get("fcff_normalizado", 0)),
        ("Valor Terminal Bruto (R$ MM)", vt.get("valor_terminal_bruto", 0)),
        ("Valor Terminal Descontado (R$ MM)", vt.get("valor_terminal_descontado", 0)),
    ]
    for i, (label, val) in enumerate(vt_rows, start=r_vt + 1):
        _cel(ws2, i, 1, label, bold=True)
        _cel(ws2, i, 2, val, accent=True)

    # ── Aba 3: DCF ────────────────────────────────────────────────────────────
    ws3 = wb.create_sheet("DCF")
    ws3.column_dimensions["A"].width = 32
    ws3.column_dimensions["B"].width = 22

    _hdr(ws3, 1, 1, "Item")
    _hdr(ws3, 1, 2, "Valor")

    dcf_rows = [
        ("Enterprise Value (R$ MM)", fmt_brl(res.get("enterprise_value"))),
        ("(-) Dívida Líquida (R$ MM)", fmt_brl(res.get("divida_liquida"))),
        ("Equity Value (R$ MM)", fmt_brl(res.get("equity_value"))),
        ("Número de Ações (MM)", fmt_2(res.get("num_acoes"))),
        ("", ""),
        ("VALOR JUSTO POR AÇÃO", f"R$ {res.get('valor_justo', 0):.2f}"),
        ("Cotação Atual", f"R$ {res.get('cotacao_atual', 0):.2f}"),
        ("Upside / Downside", f"{res.get('upside_pct', 0):+.1f}%"),
    ]
    for i, (label, val) in enumerate(dcf_rows, start=2):
        bold = label in ("VALOR JUSTO POR AÇÃO", "Upside / Downside")
        _cel(ws3, i, 1, label, bold=bold)
        _cel(ws3, i, 2, val, bold=bold, accent=bold)

    # ── Aba 4: Sensibilidade ──────────────────────────────────────────────────
    ws4 = wb.create_sheet("Sensibilidade")
    ws4.column_dimensions["A"].width = 20
    for col in range(2, 5):
        ws4.column_dimensions[get_column_letter(col)].width = 16

    ws4.cell(row=1, column=1, value=f"Valor Justo (R$) — {ticker}").font = Font(bold=True, size=12)
    _hdr(ws4, 2, 1, "WACC \\ g")
    _hdr(ws4, 2, 2, "g = 3,5%")
    _hdr(ws4, 2, 3, "g = 4,0%")
    _hdr(ws4, 2, 4, "g = 4,5%")

    wacc = p.get("wacc", 0.15)
    linhas_sens = [
        (f"WACC {(wacc-0.01)*100:.1f}%", "wacc_menos1"),
        (f"WACC {wacc*100:.1f}% (base)", "wacc_base"),
        (f"WACC {(wacc+0.01)*100:.1f}%", "wacc_mais1"),
    ]
    gs = ["3p5", "4p0", "4p5"]
    for row_i, (label, wacc_key) in enumerate(linhas_sens, start=3):
        bold = "base" in wacc_key
        _cel(ws4, row_i, 1, label, bold=bold)
        for col_i, g_key in enumerate(gs, start=2):
            key = f"g_{g_key}_{wacc_key}"
            val = sens.get(key, 0)
            _cel(ws4, row_i, col_i, f"R$ {val:.2f}" if isinstance(val, (int, float)) else val,
                 bold=bold, accent=bold)

    wb.save(saida)


# ─── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Model Builder DCF — SBWAA")
    parser.add_argument("ticker", help="Ticker do ativo (ex: PETR4)")
    parser.add_argument("--tipo", choices=["acao", "fii", "etf"], default="acao",
                        help="Tipo do ativo (default: acao)")
    args = parser.parse_args()

    ticker = args.ticker.upper().strip()
    tipo = args.tipo
    hoje = datetime.now().strftime("%Y-%m-%d")

    print(f"Model Builder (DCF) — {ticker} | {hoje}")
    print("=" * 50)

    if tipo == "etf":
        print("DCF não aplicável para ETFs. Use análise de composição e tracking error.")
        sys.exit(0)

    from fetch_fundamentals import buscar_ticker as brapi_buscar

    print(f"Carregando dados de {ticker}...")
    try:
        dados_brapi = brapi_buscar(ticker)
    except Exception as e:
        print(f"ERRO ao buscar dados Brapi de {ticker}: {e}")
        sys.exit(1)

    earnings_ctx = carregar_earnings_mais_recente(ticker)
    macro_ctx = carregar_market_researcher_hoje(hoje)

    # RAG: buscar premissas e benchmarks de valuation do setor
    contexto_rag = ""
    try:
        sys.path.insert(0, str(PROJECT_ROOT))
        from knowledge.retriever import buscar, formatar_contexto_para_agente, base_disponivel
        if base_disponivel():
            setor = dados_brapi.get("setor", "")
            query_dcf = f"DCF valuation {setor} WACC custo capital beta Brasil"
            chunks = buscar(query_dcf, n_resultados=4, filtro_tipo="research")
            contexto_rag = formatar_contexto_para_agente(chunks, max_tokens=1500)
    except Exception:
        pass

    skill_content = SKILL_PATH.read_text(encoding="utf-8")

    tipo_label = "FII" if tipo == "fii" else "AÇÃO"
    metodologia_extra = "\nMETODOLOGIA: FII — usar Gordon Growth direto no DY." if tipo == "fii" else ""

    contexto = f"""TICKER: {ticker}
TIPO: {tipo_label}
DATA: {hoje}{metodologia_extra}

DADOS BRAPI:
- Nome: {dados_brapi.get('nome', 'N/D')}
- Setor: {dados_brapi.get('setor', 'N/D')}
- Cotação atual: R$ {dados_brapi.get('cotacao', 'N/D')}
- Variação dia: {dados_brapi.get('variacao_dia_pct', 'N/D')}%
- P/L: {dados_brapi.get('pl', 'N/D')}
- EV/EBITDA: {dados_brapi.get('ev_ebitda', 'N/D')}
- P/VP: {dados_brapi.get('pvp', 'N/D')}
- Dividend Yield: {dados_brapi.get('dy', 'N/D')}%
- ROE: {dados_brapi.get('roe', 'N/D')}%
- Dívida Líq./EBITDA: {dados_brapi.get('divida_liquida_ebitda', 'N/D')}
- Dividendos recentes: {dados_brapi.get('dividendos_recentes', [])}

{f"CONTEXTO EARNINGS REVIEWER:{chr(10)}{earnings_ctx}" if earnings_ctx else "EARNINGS REVIEWER: Não disponível para este ticker."}

{f"CONTEXTO MARKET RESEARCHER:{chr(10)}{macro_ctx}" if macro_ctx else "MARKET RESEARCHER: Não disponível para hoje."}
{chr(10) + contexto_rag if contexto_rag else ""}
Construa o modelo DCF completo e retorne EXCLUSIVAMENTE o JSON estruturado.
"""

    import anthropic
    print("Enviando ao Model Builder (claude-opus-4-6)...")
    client = anthropic.Anthropic()
    response = client.messages.create(
        model="claude-opus-4-6",
        max_tokens=4000,
        system=skill_content,
        messages=[{"role": "user", "content": contexto}],
    )

    texto_resposta = response.content[0].text

    try:
        dcf = extrair_json(texto_resposta)
    except Exception as e:
        print(f"ERRO ao parsear JSON do modelo: {e}")
        print("Resposta recebida:")
        print(texto_resposta[:500])
        sys.exit(1)

    dcf["ticker"] = ticker
    dcf["data_modelo"] = hoje

    # Salvar JSON no cache
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cache_path = CACHE_DIR / f"dcf_{ticker}_{hoje}.json"
    cache_path.write_text(json.dumps(dcf, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"JSON DCF salvo: {cache_path}")

    # Gerar XLSX
    ativo_dir = VAULT_ROOT / "01-ativos" / ticker
    ativo_dir.mkdir(parents=True, exist_ok=True)

    # Versionar: dcf-v1.xlsx, dcf-v2.xlsx, ...
    versao = 1
    while (ativo_dir / f"dcf-{ticker}-v{versao}.xlsx").exists():
        versao += 1
    xlsx_path = ativo_dir / f"dcf-{ticker}-v{versao}.xlsx"
    dcf["versao"] = f"v{versao}"

    criar_xlsx(dcf, ticker, xlsx_path)
    print(f"XLSX gerado: {xlsx_path}")

    res = dcf.get("resultado", {})
    prem = dcf.get("premissas", {})
    wacc_v = prem.get("wacc", 0)
    g_v = prem.get("g_perpetuidade", 0)
    vj = res.get("valor_justo", 0)
    co = res.get("cotacao_atual", 0)
    up = res.get("upside_pct", 0)

    print(f"\n{'═'*39}")
    print(f"  DCF — {ticker}")
    print(f"{'─'*39}")
    print(f"  WACC:              {wacc_v*100:.1f}%")
    print(f"  g (perpetuidade):  {g_v*100:.1f}%")
    print(f"  Valor Justo:       R$ {vj:.2f}")
    print(f"  Cotação Atual:     R$ {co:.2f}")
    print(f"  Upside/Down:       {up:+.1f}%")
    print(f"{'═'*39}\n")

    notas = dcf.get("notas", [])
    if notas:
        print("Notas do modelo:")
        for n in notas:
            print(f"  • {n}")


if __name__ == "__main__":
    main()
